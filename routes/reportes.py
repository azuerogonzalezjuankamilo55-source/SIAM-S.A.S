import io
import logging
from datetime import date, timedelta
from typing import Any

from flask import Blueprint, flash, redirect, render_template, request, send_file, url_for
from flask_login import login_required
from decorators import admin_required
from services.reporte_service import ReporteService

logger = logging.getLogger("siam.routes.reportes")
reportes_bp = Blueprint("reportes", __name__, url_prefix="/reportes")


def _filtros():
    desde_raw = request.args.get("desde") or ""
    hasta_raw = request.args.get("hasta") or ""
    try:
        desde = date.fromisoformat(desde_raw) if desde_raw else date.today() - timedelta(days=30)
    except ValueError:
        desde = date.today() - timedelta(days=30)
    try:
        hasta = date.fromisoformat(hasta_raw) if hasta_raw else date.today()
    except ValueError:
        hasta = date.today()
    return desde, hasta


@reportes_bp.route("/")
@login_required
@admin_required
def index() -> Any:
    desde, hasta = _filtros()
    data = ReporteService.build_resumen(desde, hasta)
    return render_template("reportes/index.html", data=data, desde=desde, hasta=hasta)


@reportes_bp.route("/excel")
@login_required
@admin_required
def excel() -> Any:
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill

    desde, hasta = _filtros()
    data = ReporteService.build_resumen(desde, hasta)

    wb = Workbook()
    ws = wb.active
    ws.title = "Resumen"
    ws.append(["SIAM - Reporte de indicadores"])
    ws.append([f"Período: {desde.isoformat()} a {hasta.isoformat()}"])
    ws.append([])
    for k, v in [
        ("Facturas en el período", data.facturas_periodo),
        ("Ingresos", data.ingresos),
        ("Cartera pendiente", data.cartera),
        ("Ticket promedio", data.ticket_promedio),
        ("OTs entregadas", data.ots_entregadas),
        ("Valor del inventario", data.valor_inventario),
    ]:
        ws.append([k, v])
    for cell in ws[1]:
        cell.font = Font(bold=True)

    ws2 = wb.create_sheet("Ingresos por día")
    ws2.append(["Fecha", "Ingresos"])
    for dia, total in data.ingresos_por_dia:
        ws2.append([dia, total])

    ws3 = wb.create_sheet("Métodos de pago")
    ws3.append(["Método", "Ingresos"])
    for metodo, total in data.ingresos_por_metodo:
        ws3.append([metodo, total])

    ws4 = wb.create_sheet("Clientes")
    ws4.append(["Cliente", "Facturas", "Total"])
    for fila in data.facturacion_clientes:
        ws4.append([fila["cliente"], fila["facturas"], fila["total"]])

    ws5 = wb.create_sheet("Servicios")
    ws5.append(["Servicio", "Cantidad", "Total"])
    for fila in data.servicios_vendidos:
        ws5.append([fila["servicio"], fila["cantidad"], fila["total"]])

    ws6 = wb.create_sheet("Mecánicos")
    ws6.append(["Mecánico", "OTs entregadas"])
    for fila in data.productividad_mecanicos:
        ws6.append([fila["mecanico"], fila["ots"]])

    ws7 = wb.create_sheet("Inventario")
    ws7.append(["Tipo de movimiento", "Unidades"])
    ws7.append(["Entradas", data.movimientos_inventario["entradas"]])
    ws7.append(["Salidas", data.movimientos_inventario["salidas"]])
    ws7.append(["Ajustes", data.movimientos_inventario["ajustes"]])
    ws7.append(["Bajas", data.movimientos_inventario["bajas"]])

    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    nombre = f"reporte_siam_{desde.isoformat()}_{hasta.isoformat()}.xlsx"
    return send_file(buf, as_attachment=True, download_name=nombre,
                     mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")


@reportes_bp.route("/pdf")
@login_required
@admin_required
def pdf() -> Any:
    try:
        from weasyprint import HTML
    except ImportError:
        flash("PDF no disponible: weasyprint no está instalado", "danger")
        return redirect(url_for("reportes.index"))

    desde, hasta = _filtros()
    data = ReporteService.build_resumen(desde, hasta)
    try:
        html_str = render_template("reportes/pdf.html", data=data, desde=desde, hasta=hasta)
        pdf_bytes = HTML(string=html_str).write_pdf()
    except Exception:
        logger.error("Error generando PDF de reporte", exc_info=True)
        flash("Error generando PDF. Verifica que los binarios de WeasyPrint estén instalados.", "warning")
        return redirect(url_for("reportes.index"))

    nombre = f"reporte_siam_{desde.isoformat()}_{hasta.isoformat()}.pdf"
    return send_file(io.BytesIO(pdf_bytes), as_attachment=True, download_name=nombre,
                     mimetype="application/pdf")
